"""Phase 5 - turn your answers in review/review_queue.xlsx into config/decisions.csv.

  python scripts/apply_review.py

Reads the yellow decide_* columns on the Groups sheet (one decision per payee or
person) and the Transactions sheet (one decision per row; these win over the group).
Checks every answer against the allowed lists, appends/updates config/decisions.csv
(D0001 ...), then moves the filled-in file to review/archive/ so the next
classify.py run can build a fresh list. Invalid answers are listed and not applied.
Next: python scripts/run_all.py
"""
from __future__ import annotations

import datetime as dt
import shutil
import sys

from common import BP_VALUES, TYPES, p, read_csv, write_csv

DEC_FIELDS = ["decision_id", "scope", "key", "type", "category", "business_personal", "business_pct",
              "counterparty", "ask_accountant", "note", "decided_on", "source"]
QUEUE = p("review", "review_queue.xlsx")


def main():
    if not QUEUE.exists():
        print("No review/review_queue.xlsx - run classify.py first.")
        return
    from openpyxl import load_workbook
    wb = load_workbook(QUEUE, data_only=True)
    cats = {c["category"] for c in read_csv(p("config", "categories.csv"), required=True)}
    existing = read_csv(p("config", "decisions.csv"))
    index = {(d["scope"], d["key"]): d for d in existing}
    next_n = max((int(d["decision_id"][1:]) for d in existing if d["decision_id"][1:].isdigit()), default=0) + 1
    today = dt.date.today().isoformat()
    added = updated = 0
    errors = []

    def val(row, header, name):
        v = row[header.index(name)] if name in header else None
        return "" if v is None else str(v).strip()

    for sheet, scope_col, key_col in (("Groups", "scope", "key"), ("Transactions", None, "txn_id")):
        if sheet not in wb.sheetnames:
            continue
        rows = list(wb[sheet].iter_rows(values_only=True))
        if not rows:
            continue
        header = [str(h) if h is not None else "" for h in rows[0]]
        for r in rows[1:]:
            t = val(r, header, "decide_type")
            cat = val(r, header, "decide_category")
            bp = val(r, header, "decide_business_personal")
            pct = val(r, header, "decide_business_pct")
            person = val(r, header, "decide_person")
            ask = val(r, header, "decide_ask_accountant").upper()
            note = val(r, header, "decide_note")
            if not any([t, cat, bp, pct, person, ask, note]):
                continue
            scope = val(r, header, scope_col) if scope_col else "txn"
            key = val(r, header, key_col)
            where = f"{sheet} {key}"
            problems = []
            if t and t not in TYPES + ["Loan"]:
                problems.append(f"type {t!r} not allowed")
            if cat and cat not in cats:
                problems.append(f"category {cat!r} not in categories.csv")
            if bp and bp not in BP_VALUES:
                problems.append("business_personal must be Business, Personal or Mixed")
            if pct:
                try:
                    f = float(pct)  # a percentage, not money
                    if not 0 <= f <= 100:
                        raise ValueError
                    pct = f"{f:g}"
                except ValueError:
                    problems.append("business % must be 0-100")
            if bp == "Mixed" and not pct:
                problems.append("Mixed needs a business %")
            if not t and ask != "Y":
                problems.append("choose a type (or Y in ask_accountant)")
            if problems:
                errors.append({"where": where, "problem": "; ".join(problems)})
                continue
            d = index.get((scope, key))
            if d is None:
                d = {"decision_id": f"D{next_n:04d}", "scope": scope, "key": key}
                next_n += 1
                existing.append(d)
                index[(scope, key)] = d
                added += 1
            else:
                updated += 1
            d.update({"type": t or d.get("type", "") or ("Unknown" if ask == "Y" else ""),
                      "category": cat, "business_personal": bp, "business_pct": pct,
                      "counterparty": person, "ask_accountant": "Y" if ask == "Y" else "",
                      "note": note, "decided_on": today, "source": f"review_queue {sheet}"})

    write_csv(p("config", "decisions.csv"), existing, DEC_FIELDS)
    write_csv(p("review", "apply_review_errors.csv"), errors, ["where", "problem"])
    print(f"Decisions: {added} new, {updated} updated, {len(errors)} not applied")
    for e in errors[:10]:
        print(f"  {e['where']}: {e['problem']}")
    if errors:
        print("Fix those rows in review_queue.xlsx and run apply_review.py again (nothing was archived).")
        sys.exit(1)
    arch = p("review", "archive")
    arch.mkdir(parents=True, exist_ok=True)
    dest = arch / f"review_queue_{dt.datetime.now():%Y%m%d_%H%M%S}.xlsx"
    shutil.move(str(QUEUE), dest)
    print(f"Archived your filled-in list to {dest.relative_to(p())}")
    print("Next: python scripts/run_all.py")


if __name__ == "__main__":
    main()
