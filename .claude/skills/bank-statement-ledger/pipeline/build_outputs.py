"""Phase 6.1 - build the accountant pack.

  python scripts/build_outputs.py

Control checks run first. If A or B fails, nothing is written, so a broken pack
can't be sent by mistake:
  A. the ledger has exactly the rows combine.py produced, and each account's net
     movement per FY equals 03_data/control_totals.csv (nothing lost or doubled
     since the statements balanced)
  B. every row has one allowed type, and no txn_id appears twice
  C. matched transfer pairs net to $0.00 within each FY (pairs crossing 30 June and
     transfer rows with no matched other side are listed as open questions)
Writes:
  output/master_ledger.csv   every row, all years
  output/FY2024.xlsx ...     one workbook per financial year with the tabs README, Ledger,
                             Business P&L, Personal, Mixed use, Internal transfers, Loans,
                             Reconciliation, Open questions, Not in the bank
Reads the owner's income_not_in_bank.csv (cash kept, swaps) if present and copies it,
per FY, to the "Not in the bank" tab. Those amounts are the owner's own figures: they
never enter the Ledger or any total. A row with an unreadable date, kind or amount
stops the build.
Totals are live SUMIFS formulas over the Ledger tab (change a category and they update).
Next to each is the value this script calculated, and a Check column that is 0 while
they agree - a visible proof the formulas and the data match.
"""
from __future__ import annotations

import datetime as dt
import json
import sys
from decimal import Decimal

from common import (LOAN_TYPES, PNL_TYPES, TRANSFER_TYPES, TYPES, fmt_money, fy_bounds, fy_of, iso,
                    load_accounts, load_project, mask, money, p, parse_money, read_csv, write_csv)

ZERO = Decimal("0.00")
MONEY_FMT = '#,##0.00;[Red]-#,##0.00'
DATE_FMT = "dd/mm/yyyy"
LEDGER_COLS = ["txn_id", "date", "account_id", "account_name", "original_description", "cleaned_description",
               "counterparty", "amount", "money_in", "money_out", "type", "category", "ato_label",
               "business_personal", "business_pct", "match_id", "review_status", "notes", "classified_by",
               "confidence", "source_file", "source_page", "source_line", "balance_shown"]
ATO_ORDER = ["Purchases and other costs", "Contractor, sub-contractor and commission expenses",
             "Superannuation expenses", "Bad debts", "Lease expenses", "Rent expenses",
             "Interest expenses within Australia", "Interest expenses overseas", "Depreciation expenses",
             "Motor vehicle expenses", "Repairs and maintenance", "All other expenses"]
TYPE_HELP = {
    "Business income": "Money earned by the business (in the Business P&L).",
    "Other income": "Interest, wages, government payments - taxable, but not business income.",
    "Expense": "Spending. Business, Personal or Mixed per the business_personal column.",
    "Refund": "Money back from a supplier; reduces the category it refunds.",
    "Internal transfer": "Between two of the owner's accounts. Not income or spending.",
    "Owner drawing": "Business account -> personal account. Not an expense.",
    "Owner contribution": "Personal account -> business account. Not income.",
    "Loan out": "Money lent to a friend/family member (or repaid to them). See the Loans tab.",
    "Loan in": "Money borrowed from / repaid by a friend or family member. See the Loans tab.",
    "Gift or shared bill": "Personal money to/from people that is not a loan.",
    "Tax": "Payments to / refunds from the ATO. Not income or an expense.",
    "Outside account": "To/from an account whose statements are not in this set.",
    "Unknown": "Not classified - listed under Open questions.",
}
NOT_IN_BANK = "income_not_in_bank.csv"
NIB_KINDS = ("Cash", "Swap", "Other")


def fail(msg):
    print(f"FAIL: {msg}")
    print("Nothing was written. Fix the cause and re-run.")
    sys.exit(1)


def load_not_in_bank(fsm: int) -> list[dict] | None:
    """The owner's list of income that never went through a bank account. None if there is no file."""
    f = p(NOT_IN_BANK)
    if not f.exists():
        return None
    out = []
    for n, r in enumerate(read_csv(f), 2):
        if not any(r.values()):
            continue
        d = r.get("date", "")
        try:
            day = dt.date.fromisoformat(d if len(d) == 10 else d + "-01")
        except ValueError:
            fail(f"{NOT_IN_BANK} line {n}: date {d!r} - use YYYY-MM-DD, or YYYY-MM if you only know the month")
        kind = r.get("kind", "").capitalize()
        if kind not in NIB_KINDS:
            fail(f"{NOT_IN_BANK} line {n}: kind {r.get('kind', '')!r} - use one of {', '.join(NIB_KINDS)}")
        try:
            amt = parse_money(r.get("amount"))
        except ValueError:
            fail(f"{NOT_IN_BANK} line {n}: amount {r.get('amount')!r} is not a number (leave it blank if unknown)")
        out.append({"date": d, "_day": day, "fy": fy_of(day, fsm), "kind": kind, "amount": amt,
                    **{k: mask(r.get(k, "")) for k in ("who", "what_for", "how_you_know", "note")}})
    return out


def main():
    cfg = load_project()
    fsm = int(cfg["fy_start_month"])
    accounts = load_accounts()
    cats = {c["category"]: c for c in read_csv(p("config", "categories.csv"), required=True) if c.get("category")}
    people = {r["person"]: r for r in read_csv(p("config", "people.csv")) if r.get("person")}
    ledger = read_csv(p("03_data", "classified.csv"), required=True)
    for x in ledger:
        x["amount"] = money(x["amount"])
        x["_date"] = iso(x["date"])
    tx_count = len(read_csv(p("03_data", "transactions.csv"), required=True))
    control = {(r["account_id"], r["fy"]): (int(r["rows"]), money(r["net"]))
               for r in read_csv(p("03_data", "control_totals.csv"), required=True)}
    matches = read_csv(p("03_data", "transfer_matches.csv"))
    recon = read_csv(p("output", "reconciliation.csv"), required=True)
    unmatched = read_csv(p("review", "transfers_unmatched.csv"))
    clog_f = p("03_data", "combine_log.json")
    clog = json.loads(clog_f.read_text()) if clog_f.exists() else {"excluded_not_balanced": []}

    # ---------------------------------------------------------------- checks
    checks = []
    calc: dict[tuple, list] = {}
    for x in ledger:
        c = calc.setdefault((x["account_id"], x["fy"]), [0, ZERO])
        c[0] += 1
        c[1] += x["amount"]
    calc_t = {k: (v[0], v[1]) for k, v in calc.items()}
    if len(ledger) != tx_count:
        fail(f"check A: ledger has {len(ledger)} rows but transactions.csv has {tx_count} - re-run classify.py")
    if calc_t != control:
        diffs = [f"{k[0]} {k[1]}: ledger {v} vs statements {control.get(k)}" for k, v in calc_t.items() if control.get(k) != v]
        fail("check A: account totals do not tie to the balanced statements: " + "; ".join(diffs[:5]))
    checks.append(("A", "Every row from the balanced statements is in the ledger exactly once; "
                   "each account's net movement per FY ties to the statements", "PASS"))
    ids = [x["txn_id"] for x in ledger]
    bad_types = [x["txn_id"] for x in ledger if x["type"] not in TYPES]
    if len(ids) != len(set(ids)) or bad_types:
        fail(f"check B: duplicate txn_ids or bad types ({bad_types[:5]})")
    checks.append(("B", "Every row has exactly one allowed type; no txn_id appears twice", "PASS"))

    by_id = {x["txn_id"]: x for x in ledger}
    partner_of = {}
    for m in matches:
        partner_of[m["out_txn_id"]], partner_of[m["in_txn_id"]] = m["in_txn_id"], m["out_txn_id"]
    pair_problems = []
    for m in matches:
        o, i = by_id.get(m["out_txn_id"]), by_id.get(m["in_txn_id"])
        if not o or not i or o["type"] not in TRANSFER_TYPES or i["type"] not in TRANSFER_TYPES:
            pair_problems.append(m["match_id"])
    fy_net_problems = []
    for fy in sorted({x["fy"] for x in ledger}):
        s = sum((x["amount"] for x in ledger if x["fy"] == fy and x["type"] in TRANSFER_TYPES and x["match_id"]
                 and by_id.get(partner_of.get(x["txn_id"]), {}).get("fy") == fy), ZERO)
        if s != 0:
            fy_net_problems.append(f"{fy}: {fmt_money(s)}")
    nib = load_not_in_bank(fsm)
    checks.append(("C", "Matched transfer pairs net to $0.00 within each FY",
                   "PASS" if not fy_net_problems and not pair_problems else
                   f"CHECK: {'; '.join(fy_net_problems)} {'overridden pairs: ' + ', '.join(pair_problems[:5]) if pair_problems else ''}"))

    # ---------------------------------------------------------------- write
    master = sorted(ledger, key=lambda x: (x["_date"], x["account_id"], x["txn_id"]))
    from classify import LEDGER_FIELDS
    write_csv(p("output", "master_ledger.csv"), master, LEDGER_FIELDS)
    fys = sorted({x["fy"] for x in ledger})
    for fy in fys:
        build_fy(fy, cfg, fsm, accounts, cats, people, [x for x in master if x["fy"] == fy], master,
                 matches, recon, unmatched, clog, checks, by_id, nib)
    stray = sorted({x["fy"] for x in nib or []} - set(fys))
    if stray:
        print(f"Note: {NOT_IN_BANK} has items in {', '.join(stray)}, which has no ledger - they are in no workbook")
    print(f"Checks: " + "; ".join(f"{c[0]} {c[2]}" for c in checks))
    print(f"Wrote output/master_ledger.csv ({len(master)} rows) and " + ", ".join(f"output/{f}.xlsx" for f in fys))


def build_fy(fy, cfg, fsm, accounts, cats, people, rows, master, matches, recon, unmatched, clog, checks, by_id, nib):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    start, end = fy_bounds(fy, fsm)
    wb = Workbook()
    bold, head_fill = Font(bold=True), PatternFill("solid", fgColor="DDEBF7")
    big = Font(bold=True, size=14)

    def header(ws, r, values):
        for c, v in enumerate(values, 1):
            cell = ws.cell(r, c, v)
            cell.font = bold
            cell.fill = head_fill

    def widths(ws, w):
        for i, n in enumerate(w, 1):
            ws.column_dimensions[get_column_letter(i)].width = n

    # ---- Ledger
    led = wb.active
    led.title = "Ledger"
    header(led, 1, LEDGER_COLS)
    for x in rows:
        vals = []
        for c in LEDGER_COLS:
            v = x.get(c, "")
            if c == "date":
                v = x["_date"]
            elif c in {"amount", "money_in", "money_out", "balance_shown"}:
                v = money(v) if not isinstance(v, Decimal) else v
            elif c == "business_pct":
                v = Decimal(v) if v else None
            vals.append(v)
        led.append(vals)
    last = len(rows) + 1
    col = {c: get_column_letter(i) for i, c in enumerate(LEDGER_COLS, 1)}
    for r in range(2, last + 1):
        led[f"{col['date']}{r}"].number_format = DATE_FMT
        for c in ("amount", "money_in", "money_out", "balance_shown"):
            led[f"{col[c]}{r}"].number_format = MONEY_FMT
    led.freeze_panes = "A2"
    led.auto_filter.ref = f"A1:{get_column_letter(len(LEDGER_COLS))}{last}"
    widths(led, [12, 11, 16, 14, 42, 28, 18, 12, 11, 11, 18, 26, 30, 10, 8, 8, 14, 30, 12, 9, 34, 6, 6, 12])

    def rng(c):
        return f"Ledger!${col[c]}$2:${col[c]}${last}"
    AMT, TYP, CAT, BP, ACC = rng("amount"), rng("type"), rng("category"), rng("business_personal"), rng("account_id")

    def tot(pred):
        return sum((x["amount"] for x in rows if pred(x)), ZERO)

    # ---- Business P&L
    pl = wb.create_sheet("Business P&L")
    pl["A1"] = f"Business income and expenses {fy} ({start:%d %b %Y} - {end:%d %b %Y})"
    pl["A1"].font = big
    pl["A2"] = ("Business rows only. Mixed-use items are on the 'Mixed use' tab and are NOT in these totals. "
                "Amount = live formula over the Ledger tab; Built value = what the build script calculated; "
                "Check is 0 while they agree.")
    pl["A2"].alignment = Alignment(wrap_text=True)
    pl.merge_cells("A2:E2")
    pl.row_dimensions[2].height = 45
    header(pl, 4, ["ATO label", "Category", "Amount", "Built value", "Check (0)"])
    r = 5
    refs = {}

    def line(label, cat, formula, built):
        nonlocal r
        pl.cell(r, 1, label)
        pl.cell(r, 2, cat)
        pl.cell(r, 3, formula).number_format = MONEY_FMT
        pl.cell(r, 4, built).number_format = MONEY_FMT
        pl.cell(r, 5, f"=C{r}-D{r}").number_format = MONEY_FMT
        r += 1
        return r - 1

    pl.cell(r, 1, "INCOME").font = bold
    r += 1
    inc_cats = sorted({x["category"] for x in rows if x["type"] == "Business income" and x["business_personal"] == "Business"})
    first = r
    for c in inc_cats:
        line(cats.get(c, {}).get("ato_label", ""), c,
             f'=SUMIFS({AMT},{CAT},$B{r},{BP},"Business",{TYP},"Business income")',
             tot(lambda x, c=c: x["type"] == "Business income" and x["business_personal"] == "Business" and x["category"] == c))
    inc_row = line("Total business income", "", f"=SUM(C{first}:C{r - 1})" if inc_cats else 0,
                   tot(lambda x: x["type"] == "Business income" and x["business_personal"] == "Business"))
    pl.cell(inc_row, 1).font = bold
    r += 1
    pl.cell(r, 1, "EXPENSES (shown as positive numbers)").font = bold
    r += 1
    exp_cats = {x["category"] for x in rows if x["type"] in {"Expense", "Refund"} and x["business_personal"] == "Business"}
    labels = sorted({cats.get(c, {}).get("ato_label", "") or "(no ATO label)" for c in exp_cats},
                    key=lambda l: (ATO_ORDER.index(l) if l in ATO_ORDER else 99, l))
    sub_rows = []
    for lab in labels:
        s0 = r
        for c in sorted(c for c in exp_cats if (cats.get(c, {}).get("ato_label", "") or "(no ATO label)") == lab):
            line(lab, c,
                 f'=-(SUMIFS({AMT},{CAT},$B{r},{BP},"Business",{TYP},"Expense")+SUMIFS({AMT},{CAT},$B{r},{BP},"Business",{TYP},"Refund"))',
                 -tot(lambda x, c=c: x["type"] in {"Expense", "Refund"} and x["business_personal"] == "Business" and x["category"] == c))
        sr = line(f"Subtotal: {lab}", "", f"=SUM(C{s0}:C{r - 1})",
                  -tot(lambda x, lab=lab: x["type"] in {"Expense", "Refund"} and x["business_personal"] == "Business"
                       and (cats.get(x["category"], {}).get("ato_label", "") or "(no ATO label)") == lab))
        pl.cell(sr, 1).font = bold
        sub_rows.append(sr)
    exp_row = line("Total business expenses", "", "=" + "+".join(f"C{s}" for s in sub_rows) if sub_rows else 0,
                   -tot(lambda x: x["type"] in {"Expense", "Refund"} and x["business_personal"] == "Business"))
    pl.cell(exp_row, 1).font = bold
    r += 1
    net_row = line("Net business result (income - expenses)", "", f"=C{inc_row}-C{exp_row}",
                   tot(lambda x: x["type"] == "Business income" and x["business_personal"] == "Business")
                   + tot(lambda x: x["type"] in {"Expense", "Refund"} and x["business_personal"] == "Business"))
    pl.cell(net_row, 1).font = bold
    refs.update(income=f"'Business P&L'!C{inc_row}", expenses=f"'Business P&L'!C{exp_row}", net=f"'Business P&L'!C{net_row}")
    widths(pl, [44, 34, 16, 16, 12])

    # ---- Mixed use
    mx = wb.create_sheet("Mixed use")
    mx["A1"] = "Mixed business/personal items - business % is a starting figure for the accountant to confirm"
    mx["A1"].font = bold
    header(mx, 3, ["txn_id", "Date", "Description", "Category", "ATO label", "Amount", "Business %", "Business portion"])
    mrows = [x for x in rows if x["business_personal"] == "Mixed" and x["type"] in PNL_TYPES]
    r = 4
    for x in mrows:
        mx.append([x["txn_id"], x["_date"], x["cleaned_description"], x["category"], x["ato_label"], x["amount"],
                   Decimal(x["business_pct"]) if x["business_pct"] else None, f"=-F{r}*G{r}/100"])
        mx[f"B{r}"].number_format = DATE_FMT
        mx[f"F{r}"].number_format = MONEY_FMT
        mx[f"H{r}"].number_format = MONEY_FMT
        r += 1
    mx.cell(r, 3, "Total business portion (expenses positive)").font = bold
    mx.cell(r, 8, f"=SUM(H4:H{r - 1})" if mrows else 0).number_format = MONEY_FMT
    refs["mixed"] = f"'Mixed use'!H{r}"
    widths(mx, [12, 11, 40, 26, 30, 12, 10, 16])

    # ---- Personal
    ps = wb.create_sheet("Personal")
    ps["A1"] = f"Other income and personal spending {fy}"
    ps["A1"].font = big
    r = 3
    header(ps, r, ["Other income (taxable, not business income)", "ATO item", "Amount", "Built value", "Check (0)"])
    r += 1
    o0 = r
    for c in sorted({x["category"] for x in rows if x["type"] == "Other income"}):
        ps.cell(r, 1, c)
        ps.cell(r, 2, cats.get(c, {}).get("ato_label", ""))
        ps.cell(r, 3, f'=SUMIFS({AMT},{CAT},$A{r},{TYP},"Other income")').number_format = MONEY_FMT
        ps.cell(r, 4, tot(lambda x, c=c: x["type"] == "Other income" and x["category"] == c)).number_format = MONEY_FMT
        ps.cell(r, 5, f"=C{r}-D{r}").number_format = MONEY_FMT
        r += 1
    ps.cell(r, 1, "Total other income").font = bold
    ps.cell(r, 3, f"=SUM(C{o0}:C{r - 1})" if r > o0 else 0).number_format = MONEY_FMT
    refs["other_income"] = f"Personal!C{r}"
    r += 2
    header(ps, r, ["Interest earned by account", "Account", "Amount", "Built value", "Check (0)"])
    r += 1
    i0 = r
    for aid in sorted({x["account_id"] for x in rows if x["type"] == "Other income" and x["category"] == "Interest earned"}):
        ps.cell(r, 1, "Interest earned")
        ps.cell(r, 2, aid)
        ps.cell(r, 3, f'=SUMIFS({AMT},{CAT},"Interest earned",{ACC},$B{r},{TYP},"Other income")').number_format = MONEY_FMT
        ps.cell(r, 4, tot(lambda x, aid=aid: x["type"] == "Other income" and x["category"] == "Interest earned"
                          and x["account_id"] == aid)).number_format = MONEY_FMT
        ps.cell(r, 5, f"=C{r}-D{r}").number_format = MONEY_FMT
        r += 1
    refs["interest"] = f"SUM(Personal!C{i0}:C{max(r - 1, i0)})"
    r += 1
    header(ps, r, ["Personal spending by category (positive)", "", "Amount", "Built value", "Check (0)"])
    r += 1
    for c in sorted({x["category"] for x in rows if x["type"] in {"Expense", "Refund"} and x["business_personal"] == "Personal"}):
        ps.cell(r, 1, c)
        ps.cell(r, 3, f'=-(SUMIFS({AMT},{CAT},$A{r},{BP},"Personal",{TYP},"Expense")+SUMIFS({AMT},{CAT},$A{r},{BP},"Personal",{TYP},"Refund"))').number_format = MONEY_FMT
        ps.cell(r, 4, -tot(lambda x, c=c: x["type"] in {"Expense", "Refund"} and x["business_personal"] == "Personal"
                           and x["category"] == c)).number_format = MONEY_FMT
        ps.cell(r, 5, f"=C{r}-D{r}").number_format = MONEY_FMT
        r += 1
    r += 1
    ps.cell(r, 1, "Personal items that may be deductible (accountant to confirm)").font = bold
    r += 1
    header(ps, r, ["Date", "Description", "Category", "Amount", "txn_id"])
    r += 1
    for x in rows:
        if cats.get(x["category"], {}).get("possibly_deductible", "").upper() == "Y" and x["business_personal"] == "Personal":
            ps.cell(r, 1, x["_date"]).number_format = DATE_FMT
            ps.cell(r, 2, x["cleaned_description"])
            ps.cell(r, 3, x["category"])
            ps.cell(r, 4, x["amount"]).number_format = MONEY_FMT
            ps.cell(r, 5, x["txn_id"])
            r += 1
    r += 1
    ps.cell(r, 1, "Payments to / refunds from the ATO (not income or expenses)").font = bold
    r += 1
    header(ps, r, ["Date", "Description", "Category", "Amount", "txn_id"])
    r += 1
    for x in rows:
        if x["type"] == "Tax":
            ps.cell(r, 1, x["_date"]).number_format = DATE_FMT
            ps.cell(r, 2, x["cleaned_description"])
            ps.cell(r, 3, x["category"])
            ps.cell(r, 4, x["amount"]).number_format = MONEY_FMT
            ps.cell(r, 5, x["txn_id"])
            r += 1
    widths(ps, [46, 34, 22, 16, 14])

    # ---- Internal transfers
    it = wb.create_sheet("Internal transfers")
    it["A1"] = "Money moved between the owner's own accounts - not income or spending"
    it["A1"].font = bold
    header(it, 3, ["match_id", "Type", "Out date", "From account", "In date", "To account", "Amount",
                   "Days apart", "Confidence", "Crosses 30 June", "Note"])
    r = 4
    for m in matches:
        if not (start <= iso(m["out_date"]) <= end or start <= iso(m["in_date"]) <= end):
            continue
        it.append([m["match_id"], m["type"], iso(m["out_date"]), m["out_account"], iso(m["in_date"]),
                   m["in_account"], money(m["amount"]), int(m["days_apart"]), m["confidence"], m["crosses_fy"], m["note"]])
        it[f"C{r}"].number_format = it[f"E{r}"].number_format = DATE_FMT
        it[f"G{r}"].number_format = MONEY_FMT
        r += 1
    r += 1
    it.cell(r, 1, "Totals this FY").font = bold
    for t in ("Owner drawing", "Owner contribution", "Internal transfer"):
        r += 1
        it.cell(r, 1, t)
        it.cell(r, 7, f'=SUMIFS({AMT},{TYP},"{t}",{AMT},">0")').number_format = MONEY_FMT
        it.cell(r, 8, "(money received side)")
    refs["drawings"] = f"'Internal transfers'!G{r - 2}"
    refs["contributions"] = f"'Internal transfers'!G{r - 1}"
    widths(it, [10, 18, 11, 18, 11, 18, 12, 10, 11, 10, 50])

    # ---- Loans
    ln = wb.create_sheet("Loans")
    ln["A1"] = "Informal loans with friends and family. Net owed to you: + means they owe you, - means you owe them."
    ln["A1"].font = bold
    loans_all = [x for x in master if x["type"] in LOAN_TYPES and x["_date"] <= end]
    thresh = Decimal(str(cfg["loan_flag_threshold"]))
    header(ln, 3, ["Person", "Lent / repaid by you this FY", "Received this FY", "Net this FY",
                   "Net owed to you at 30 June", "Flags"])
    r = 4
    persons = sorted({x["counterparty"] or "(unnamed)" for x in loans_all})
    for person in persons:
        mine = [x for x in master if x["type"] in LOAN_TYPES and (x["counterparty"] or "(unnamed)") == person]
        upto = [x for x in mine if x["_date"] <= end]
        this = [x for x in upto if x["_date"] >= start]
        out_ = -sum((x["amount"] for x in this if x["amount"] < 0), ZERO)
        in_ = sum((x["amount"] for x in this if x["amount"] > 0), ZERO)
        flags = []
        if mine and all(x["amount"] > 0 for x in mine):
            flags.append("only ever paid you - check it is not income")
        if any(abs(x["amount"]) >= thresh for x in mine):
            flags.append(f"an amount of ${thresh:,.0f} or more")
        if people.get(person, {}).get("relationship", "") == "customer":
            flags.append("also a business customer")
        if person == "(unnamed)":
            flags.append("person not named - use decide_person on the review list")
        ln.append([person, out_, in_, out_ - in_, -sum((x["amount"] for x in upto), ZERO), "; ".join(flags)])
        for c in "BCDE":
            ln[f"{c}{r}"].number_format = MONEY_FMT
        r += 1
    loans_total_row = r
    ln.cell(r, 1, "Total").font = bold
    ln.cell(r, 5, f"=SUM(E4:E{r - 1})" if persons else 0).number_format = MONEY_FMT
    refs["loans"] = f"Loans!E{loans_total_row}"
    r += 2
    ln.cell(r, 1, "Detail (all loan rows up to 30 June, running net per person)").font = bold
    r += 1
    header(ln, r, ["Person", "Date", "Description", "Amount", "Running net owed to you", "txn_id"])
    r += 1
    running: dict[str, Decimal] = {}
    for x in sorted(loans_all, key=lambda x: ((x["counterparty"] or "(unnamed)"), x["_date"], x["txn_id"])):
        who = x["counterparty"] or "(unnamed)"
        running[who] = running.get(who, ZERO) - x["amount"]
        ln.append([who, x["_date"], x["original_description"], x["amount"], running[who], x["txn_id"]])
        ln[f"B{r}"].number_format = DATE_FMT
        ln[f"D{r}"].number_format = ln[f"E{r}"].number_format = MONEY_FMT
        r += 1
    widths(ln, [22, 16, 42, 14, 22, 40])

    # ---- Reconciliation
    rc = wb.create_sheet("Reconciliation")
    rc["A1"] = "Balance check for every statement in this FY (opening + transactions = closing)"
    rc["A1"].font = bold
    cols = ["statement_id", "account_id", "period_start", "period_end", "rows", "opening_balance", "total_in",
            "total_out", "computed_closing", "stated_closing", "difference", "row_errors", "continuity", "status",
            "source_file", "note"]
    header(rc, 3, cols)
    r = 4
    fy_recon = [s for s in recon if s.get("period_start") and not (iso(s["period_end"]) < start or iso(s["period_start"]) > end)]
    for s in fy_recon:
        vals = []
        for c in cols:
            v = s.get(c, "")
            if c in {"opening_balance", "total_in", "total_out", "computed_closing", "stated_closing", "difference"}:
                v = money(v)
            elif c in {"period_start", "period_end"}:
                v = iso(v)
            elif c in {"rows", "row_errors"} and str(v).isdigit():
                v = int(v)
            vals.append(v)
        rc.append(vals)
        for c in "CD":
            rc[f"{c}{r}"].number_format = DATE_FMT
        for c in "FGHIJK":
            rc[f"{c}{r}"].number_format = MONEY_FMT
        r += 1
    r += 1
    rc.cell(r, 1, "Account balances (from the balanced statements)").font = bold
    r += 1
    header(rc, r, ["account_id", "Balance at FY start", "Balance at FY end", "Movement in ledger", "Note"])
    r += 1
    for aid in sorted(accounts):
        chain = sorted([s for s in recon if s["account_id"] == aid and s.get("status") in {"PASS", "SIGNED_OFF"}
                        and s.get("period_start") and not str(s.get("continuity", "")).startswith("OVERLAP")],
                       key=lambda s: s["period_start"])
        if not chain:
            continue
        opening0 = money(chain[0]["opening_balance"])
        acct_rows = [x for x in master if x["account_id"] == aid]
        note = ""
        if iso(chain[0]["period_start"]) > start:
            note = f"first statement starts {chain[0]['period_start']}"
        b_start = opening0 + sum((x["amount"] for x in acct_rows if x["_date"] < start), ZERO)
        b_end = opening0 + sum((x["amount"] for x in acct_rows if x["_date"] <= end), ZERO)
        rc.append([aid, b_start, b_end, b_end - b_start, note])
        for c in "BCD":
            rc[f"{c}{r}"].number_format = MONEY_FMT
        r += 1
    widths(rc, [12, 18, 12, 12, 7, 14, 12, 12, 14, 14, 11, 9, 30, 11, 38, 40])

    # ---- Open questions
    oq = wb.create_sheet("Open questions")
    oq["A1"] = "Open questions and exceptions for this FY"
    oq["A1"].font = bold
    header(oq, 3, ["Topic", "Reference", "Date", "Amount", "Detail"])
    items = []
    for x in rows:
        if x["review_status"] == "Needs review":
            items.append(("Not yet reviewed", x["txn_id"], x["_date"], x["amount"], f"{x['original_description']} - {x['review_reasons']}"))
        elif x["review_status"] == "Ask accountant":
            items.append(("Question for accountant", x["txn_id"], x["_date"], x["amount"], f"{x['original_description']} - {x['notes']}"))
        if "large business item" in x["notes"]:
            items.append(("Large business item", x["txn_id"], x["_date"], x["amount"], x["original_description"]))
        if x["type"] in TRANSFER_TYPES and not x["match_id"]:
            items.append(("Transfer with no matched other side", x["txn_id"], x["_date"], x["amount"], x["original_description"]))
    for u in unmatched:
        if start <= iso(u["date"]) <= end and by_id.get(u["txn_id"], {}).get("type") not in TRANSFER_TYPES | {"Outside account"}:
            items.append(("Looks like a transfer, no partner found", u["txn_id"], iso(u["date"]), money(u["amount"]),
                          f"{u['description_raw']} - {u['why_flagged']}"))
    for m in matches:
        if m["crosses_fy"] and (start <= iso(m["out_date"]) <= end or start <= iso(m["in_date"]) <= end):
            items.append(("Transfer crosses 30 June", m["match_id"], iso(m["out_date"]), money(m["amount"]),
                          f"{m['out_account']} -> {m['in_account']}, arrives {m['in_date']}"))
    for rr in range(4, loans_total_row):
        flag = ln.cell(rr, 6).value
        if flag:
            items.append(("Loan flag", ln.cell(rr, 1).value, None, ln.cell(rr, 5).value, flag))
    for s in fy_recon:
        if s.get("status") not in {"PASS", "SIGNED_OFF"}:
            items.append(("Statement does not balance", s["statement_id"], iso(s["period_start"]), money(s.get("difference", "")), s.get("note", "")))
        if str(s.get("continuity", "")).startswith(("BREAK", "GAP")):
            items.append(("Statement continuity", s["statement_id"], iso(s["period_start"]), None, s["continuity"]))
    for sid in clog.get("excluded_not_balanced", []):
        items.append(("Statement excluded (not balanced)", sid, None, None, "its rows are NOT in this ledger"))
    cov = read_csv(p("output", "coverage.csv"))
    for line_ in cov:
        miss = [m for m, v in line_.items() if m != "account_id" and v == "MISSING" and start.isoformat()[:7] <= m <= end.isoformat()[:7]]
        if miss:
            items.append(("Months with no statement", line_["account_id"], None, None, ", ".join(miss)))
    r = 4
    for it_ in items:
        oq.append(list(it_))
        oq[f"C{r}"].number_format = DATE_FMT
        oq[f"D{r}"].number_format = MONEY_FMT
        r += 1
    widths(oq, [34, 16, 11, 14, 90])

    # ---- Not in the bank (the owner's own list; never in the Ledger or a total)
    nb = wb.create_sheet("Not in the bank")
    nb["A1"] = "Income that never went through a bank account - the owner's own list (cash kept, swaps)"
    nb["A1"].font = bold
    nb["A2"] = ("Not from the bank statements, and NOT in the Ledger or any total in this pack. Amounts are the "
                "owner's figures; for a swap, the owner's rough value of what they received. The accountant decides "
                "how each item is treated.")
    header(nb, 4, ["Date", "Kind", "Who", "What for", "Amount or value", "How the owner knows", "Note"])
    mine = sorted((x for x in nib or [] if x["fy"] == fy), key=lambda x: x["_day"])
    r = 5
    for x in mine:
        nb.append([x["date"], x["kind"], x["who"], x["what_for"], x["amount"], x["how_you_know"], x["note"]])
        nb[f"E{r}"].number_format = MONEY_FMT
        r += 1
    if nib is None:
        nb.cell(5, 1, f"No list provided ({NOT_IN_BANK} is not in the project folder).")
        refs["not_in_bank"] = '"no list provided"'
    else:
        if not mine:
            nb.cell(5, 1, "The owner listed nothing for this FY.")
        end_ = max(r - 1, 5)
        for i, kind in enumerate(NIB_KINDS):
            nb.cell(end_ + 2 + i, 4, f"Total {kind.lower()}")
            c = nb.cell(end_ + 2 + i, 5, f'=SUMIF(B5:B{end_},"{kind}",E5:E{end_})')
            c.number_format = MONEY_FMT
        refs["not_in_bank"] = f"SUM('Not in the bank'!E5:E{end_})"
    widths(nb, [11, 8, 24, 36, 16, 30, 40])

    # ---- README (first tab)
    rd = wb.create_sheet("README", 0)
    needs = sum(1 for x in rows if x["review_status"] == "Needs review")
    reasons = []
    if clog.get("excluded_not_balanced"):
        reasons.append(f"{len(clog['excluded_not_balanced'])} statements excluded because they do not balance")
    if needs:
        reasons.append(f"{needs} rows not yet reviewed")
    if any(s.get("status") not in {"PASS", "SIGNED_OFF"} for s in fy_recon):
        reasons.append("some statements in this FY do not balance")
    status = "DRAFT - " + "; ".join(reasons) if reasons else "FINAL - all statements balance and every row is classified"
    rd["A1"] = f"Ledger pack {fy}: {start:%d %B %Y} to {end:%d %B %Y}"
    rd["A1"].font = Font(bold=True, size=16)
    rd["A2"] = f"Status: {status}"
    rd["A2"].font = Font(bold=True, color="C00000" if reasons else "006100")
    rd["A3"] = f"Built {dt.datetime.now():%d %b %Y %H:%M}. Currency {cfg['currency']}. " \
               f"{'GST registered - amounts are GST-inclusive.' if cfg.get('gst_registered') else 'Not registered for GST.'}"
    r = 5
    rd.cell(r, 1, "Headline figures (live links to the tabs)").font = bold
    r += 1
    for label, ref in [("Business income", refs["income"]), ("Business expenses", refs["expenses"]),
                       ("Net business result", refs["net"]),
                       ("Mixed-use items, business portion at stated % (NOT in the business result)", refs["mixed"]),
                       ("Other income (interest, wages, government payments)", refs["other_income"]),
                       ("  of which bank interest", refs["interest"]),
                       ("Owner drawings (business -> personal)", refs["drawings"]),
                       ("Owner contributions (personal -> business)", refs["contributions"]),
                       ("Informal loans: net owed to you at 30 June", refs["loans"]),
                       ("Cash kept and swaps - owner's list, NOT included above (Not in the bank tab)", refs["not_in_bank"]),
                       ("Open questions and exceptions", f"COUNTA('Open questions'!A4:A{max(len(items) + 3, 4)})")]:
        rd.cell(r, 1, label)
        c = rd.cell(r, 2, f"={ref}")
        c.number_format = MONEY_FMT if "Open questions" not in label else "0"
        r += 1
    r += 1
    rd.cell(r, 1, "Tabs").font = bold
    for name, what in [("Ledger", "Every transaction in this FY, one row each. Filter by type / business_personal / category."),
                       ("Business P&L", "Business income and expenses by ATO label and category."),
                       ("Mixed use", "Items used partly for business, with the owner's starting business %."),
                       ("Personal", "Interest and other income, personal spending, possibly-deductible personal items, ATO payments."),
                       ("Internal transfers", "Matched pairs between the owner's own accounts (excluded from income and spending)."),
                       ("Loans", "Informal loans with friends and family, net per person."),
                       ("Reconciliation", "Opening + transactions = closing for every statement; account balances."),
                       ("Open questions", "Everything not settled, for the accountant or the owner."),
                       ("Not in the bank", "The owner's own list of cash kept and swaps. Not from the statements; in no total.")]:
        r += 1
        rd.cell(r, 1, name)
        rd.cell(r, 2, what)
    r += 2
    rd.cell(r, 1, "Types").font = bold
    for t in TYPES:
        r += 1
        rd.cell(r, 1, t)
        rd.cell(r, 2, TYPE_HELP[t])
    r += 2
    rd.cell(r, 1, "How to trace any number").font = bold
    r += 1
    rd.cell(r, 1, "Each Ledger row has source_file + source_page (PDF) or source_line (CSV). txn_id S037-0042 = "
                  "statement 37, row 42. Amounts were extracted by script and every statement balances to the cent.")
    r += 2
    rd.cell(r, 1, "Control checks").font = bold
    for cid, desc, res in checks:
        r += 1
        rd.cell(r, 1, f"{cid}. {desc}")
        rd.cell(r, 2, res)
    r += 1
    rd.cell(r, 1, "D. Formula totals equal the build script's totals")
    rd.cell(r, 2, "see the Check (0) columns")
    r += 2
    rd.cell(r, 1, "Account numbers are masked to the last 4 digits. Business % figures and deductibility are the "
                  "accountant's call; this pack records the facts and the owner's starting positions.")
    widths(rd, [78, 60])

    wb.save(p("output", f"{fy}.xlsx"))


if __name__ == "__main__":
    main()
