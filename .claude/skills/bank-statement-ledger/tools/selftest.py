"""Regression test for the pipeline: builds the demo in a temp folder, runs everything,
compares the ledger to the generator's ground truth, then tries to break it on purpose.

  pip install pdfplumber openpyxl reportlab
  python <skill>/tools/selftest.py            # prints PASS/FAIL per test, exit 1 on any FAIL

Run it after changing any pipeline script. When a real project hits a new failure,
add a test here that reproduces it (and a line to the skill's Gotchas).
"""
from __future__ import annotations

import collections
import csv
import datetime as dt
import shutil
import subprocess
import sys
import tempfile
from decimal import Decimal
from pathlib import Path

TOOLS = Path(__file__).resolve().parent
sys.path.insert(0, str(TOOLS))
RESULTS = []


def check(name, ok, detail=""):
    RESULTS.append((name, ok))
    print(f"{'PASS' if ok else 'FAIL'}  {name}{'  - ' + detail if detail and not ok else ''}")


def run(root, script, *args):
    r = subprocess.run([sys.executable, str(root / "scripts" / script), *args], cwd=root,
                       capture_output=True, text=True)
    return r.returncode, r.stdout + r.stderr


def rows(path):
    with open(path, newline="", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def write_rows(path, data):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=list(data[0].keys()))
        w.writeheader()
        w.writerows(data)


def unit_tests(root):
    sys.path.insert(0, str(root / "scripts"))
    import common as c
    D = Decimal
    money_cases = {"1,234.56": D("1234.56"), "$1,234.56 CR": D("1234.56"), "12.00 DR": D("-12.00"),
                   "(45.20)": D("-45.20"), "45.20-": D("-45.20"), "-0.50": D("-0.50"), "": None, "Nil": D("0.00")}
    ok = all(c.parse_money(k) == v for k, v in money_cases.items())
    try:
        c.parse_money("12,34O.00")
        ok = False
    except ValueError:
        pass
    check("parse_money handles CR/DR, brackets, trailing minus; rejects junk", ok)
    p0, p1 = dt.date(2022, 12, 1), dt.date(2023, 1, 31)
    ok = (c.parse_date("03 Jan", p0, p1) == dt.date(2023, 1, 3) and c.parse_date("28 Dec", p0, p1) == dt.date(2022, 12, 28)
          and c.parse_date("03/04/2023") == dt.date(2023, 4, 3) and c.parse_date("3/4/23") == dt.date(2023, 4, 3)
          and c.parse_date("2023-04-03") == dt.date(2023, 4, 3))
    try:
        c.parse_date("03 Jan")
        ok = False
    except ValueError:
        pass
    check("parse_date is day-first, infers missing years from the period, refuses to guess", ok)
    m = {"BSB 062-000 12345678": "BSB xx5678", "CARD 4564 1234 5678 9012": "CARD xx9012",
         "Account number 1234 5678": "Account number xx5678", "REF 987654321": "REF xx4321",
         "01/07/2022 1,234.56 123456.78": "01/07/2022 1,234.56 123456.78", "2022-07-01": "2022-07-01"}
    bad = {k: c.mask(k) for k, v in m.items() if c.mask(k) != v}
    check("mask keeps last 4 of account/card numbers and leaves dates and amounts alone", not bad, str(bad))
    check("fy_of: 30 Jun 2023 = FY2023, 1 Jul 2023 = FY2024",
          c.fy_of(dt.date(2023, 6, 30)) == "FY2023" and c.fy_of(dt.date(2023, 7, 1)) == "FY2024")


def main():
    from make_demo import build_events

    tmp = Path(tempfile.mkdtemp(prefix="ledger-selftest-"))
    root = tmp / "books"
    subprocess.run([sys.executable, str(TOOLS / "make_demo.py"), str(root)], capture_output=True, text=True, check=True)
    unit_tests(root)

    # ---- importing from a messy folder: flatten the demo's statements, keep one file that has
    #      no account number in a named folder, add a duplicate, then sort them back into 00_raw
    src = tmp / "messy source" / "Bank Statements"
    (src / "Demo Bank" / "Everyday").mkdir(parents=True)
    original = {}
    for f in (root / "00_raw").rglob("*.*"):
        original[f.name] = f.parent.relative_to(root / "00_raw").as_posix()
        dest = src / "Demo Bank" / "Everyday" if f.name.startswith("export") else src
        shutil.copy2(f, dest / f.name)
    shutil.copy2(next((root / "00_raw" / "DemoBank" / "Business").glob("*.pdf")), src / "copy of a statement.pdf")
    (src / "notes.txt").write_text("not a statement")
    shutil.rmtree(root / "00_raw")
    (root / "00_raw").mkdir()
    code, out = run(root, "import_statements.py", str(src), "--apply")
    placed = {f.name: f.parent.relative_to(root / "00_raw").as_posix() for f in (root / "00_raw").rglob("*.*")}
    check("import_statements puts every statement in its account folder (transfer lines don't mislead it)",
          code == 0 and placed == original, f"{out[-400:]} diff={set(placed.items()) ^ set(original.items())}")
    code, out = run(root, "import_statements.py", str(src), "--apply")
    check("import_statements is safe to re-run (nothing copied twice)",
          sum(1 for _ in (root / "00_raw").rglob("*.*")) == len(original))
    run(root, "inventory.py")

    code, out = run(root, "parse_all.py")
    check("all demo statements parse", code == 0, out[-300:])
    code, out = run(root, "run_all.py")
    check("full pipeline runs on the demo", code == 0, out[-500:])

    truth = collections.Counter((a, e[0].isoformat(), e[2]) for a, es in build_events().items() for e in es)
    led = collections.Counter((r["account_id"], r["date"], Decimal(r["amount"])) for r in rows(root / "output" / "master_ledger.csv"))
    check("ledger equals ground truth (every row, date and amount; duplicates removed)", truth == led,
          f"missing {list((truth - led).items())[:3]} extra {list((led - truth).items())[:3]}")
    m = rows(root / "03_data" / "transfer_matches.csv")
    types = collections.Counter(x["type"] for x in m)
    check("transfers: 9 owner drawings + 11 internal, 1 crosses 30 June, 2 flagged ambiguous",
          types == {"Owner drawing": 9, "Internal transfer": 11} and sum(1 for x in m if x["crosses_fy"]) == 1
          and sum(1 for x in m if x["note"].startswith("ambiguous")) == 2, str(types))
    cl = rows(root / "03_data" / "classified.csv")
    cash = [x for x in cl if "CASH DEPOSIT" in x["original_description"]]
    check("cash deposit goes to review", cash and cash[0]["review_status"] == "Needs review")
    loans = {x["counterparty"] for x in cl if x["type"] in ("Loan in", "Loan out")}
    check("loans attributed to the right people", loans == {"Alex Nguyen", "Sam Lee"}, str(loans))
    ps = rows(root / "03_data" / "payees.csv")
    st = {x["payee_key"]: x["status"] for x in ps}
    check("payee status separates own transfers and people from payees needing rules",
          st.get("TRANSFER FROM XX4321 J CITIZEN") == "own-transfer" and st.get("OSKO PAYMENT TO ALEX NGUYEN") == "person"
          and st.get("BUNNINGS 7722 PENRITH") == "new")

    # ---- adversarial: a sign error in one extracted row
    x = root / "02_extracted" / "S005.csv"
    backup = x.read_text(encoding="utf-8-sig")
    data = rows(x)
    data[4]["amount"] = str(-Decimal(data[4]["amount"]))
    write_rows(x, data)
    code, out = run(root, "reconcile.py")
    errs = rows(root / "review" / "reconcile_row_errors.csv")
    check("sign error: reconcile FAILs and names the exact row", code == 1 and "S005 FAIL" in out
          and errs and errs[0]["seq"] == data[4]["seq"], out[-300:])
    code, out = run(root, "combine.py")
    check("combine refuses while a statement does not balance", code == 1)
    x.write_text(backup, encoding="utf-8-sig")

    # ---- adversarial: day/month swap
    data = rows(x)
    d = dt.date.fromisoformat(data[2]["date"])
    data[2]["date"] = dt.date(d.year, d.day, d.month).isoformat() if d.day <= 12 else "2022-01-08"
    write_rows(x, data)
    code, out = run(root, "reconcile.py")
    check("day/month swap: flagged as dates outside the period", code == 1 and "dates outside period" in out, out[-300:])
    x.write_text(backup, encoding="utf-8-sig")

    # ---- adversarial: a missing statement
    inv = root / "config" / "inventory.csv"
    inv_backup = inv.read_text(encoding="utf-8-sig")
    data = rows(inv)
    for r in data:
        if r["statement_id"] == "S005":
            r["include"] = "N"
    write_rows(inv, data)
    code, out = run(root, "reconcile.py")
    check("missing statement: continuity BREAK and MISSING months reported", code == 1 and "BREAK" in out, out[-300:])
    inv.write_text(inv_backup, encoding="utf-8-sig")

    # ---- adversarial: overlapping CSV export that disagrees with the PDF but balances on its own
    x7 = root / "02_extracted" / "S007.csv"
    b7 = x7.read_text(encoding="utf-8-sig")
    data = rows(x7)
    delta = Decimal("1.00")
    data[3]["amount"] = str(Decimal(data[3]["amount"]) - delta)
    for r in data[3:]:
        r["balance_shown"] = str(Decimal(r["balance_shown"]) - delta)
    write_rows(x7, data)
    meta7 = root / "02_extracted" / "S007.meta.json"
    mb = meta7.read_text()
    import json
    mj = json.loads(mb)
    mj["closing_balance"] = str(Decimal(mj["closing_balance"]) - delta)
    meta7.write_text(json.dumps(mj))
    code, out = run(root, "reconcile.py")
    code2, out2 = run(root, "combine.py")
    check("overlapping sources that disagree: combine stops with overlap conflicts",
          code == 0 and code2 == 1 and "overlapping" in out2, out2[-300:])
    x7.write_text(b7, encoding="utf-8-sig")
    meta7.write_text(mb)
    run(root, "run_all.py")

    # ---- owner review round trip
    from openpyxl import load_workbook
    q = root / "review" / "review_queue.xlsx"
    wb = load_workbook(q)
    ws = wb["Groups"]
    head = [c.value for c in ws[1]]
    col = {h: i + 1 for i, h in enumerate(head)}
    answers = {"BUNNINGS 7722 PENRITH": ("Expense", "Shopping and household", "Personal"),
               "DAN MURPHYS 555 PENRITH": ("Expense", "Dining and takeaway", "Personal"),
               "CASH DEPOSIT BRANCH PARRAMATTA": ("Business income", "Sales and fees", "Business"),
               "JB HI FI 0123 PARRAMATTA": ("Expense", "Equipment purchases", "Business"),
               "BPAY ATO TAX OFFICE PAYMENTS": ("Tax", "Tax payment", "Personal"),
               "INTERNET TRANSFER TO XX3456": ("Internal transfer", "", ""),
               "TRANSFER FROM J CITIZEN": ("Internal transfer", "", "")}
    for r in range(2, ws.max_row + 1):
        key = ws.cell(r, col["key"]).value
        if key in answers:
            t, cat, bp = answers[key]
            ws.cell(r, col["decide_type"], t)
            ws.cell(r, col["decide_category"], cat or None)
            ws.cell(r, col["decide_business_personal"], bp or None)
        elif key and key.startswith("AMAZON"):
            ws.cell(r, col["decide_ask_accountant"], "Y")
            ws.cell(r, col["decide_note"], "work laptop bag or personal?")
    wb.save(q)
    code, out = run(root, "classify.py")
    check("classify refuses to overwrite a filled-in review list", code == 1 and "apply_review" in out)
    code, out = run(root, "run_all.py")
    cl = rows(root / "03_data" / "classified.csv")
    open_ = [x for x in cl if x["review_status"] == "Needs review"]
    check("after the owner's answers nothing is left to review", code == 0 and not open_,
          f"{len(open_)} left: {[x['payee_key'] for x in open_][:5]} {out[-300:]}")
    amb = [x for x in cl if x["match_id"] and "+D" in x["classified_by"]]
    check("confirmed ambiguous transfers keep their pair", len(amb) >= 4, str(len(amb)))
    wb = load_workbook(root / "output" / "FY2023.xlsx")
    check("FY2023 pack status is FINAL after review", str(wb["README"]["A2"].value).startswith("Status: FINAL"),
          wb["README"]["A2"].value)

    # ---- tampering after classification must block the build
    cf = root / "03_data" / "classified.csv"
    data = rows(cf)
    write_rows(cf, data[:-1])
    code, out = run(root, "build_outputs.py")
    check("a row lost after classification blocks the build (check A)", code == 1 and "check A" in out, out[-200:])

    ok = all(r[1] for r in RESULTS)
    print(f"\n{sum(1 for r in RESULTS if r[1])}/{len(RESULTS)} passed")
    if ok:
        shutil.rmtree(tmp, ignore_errors=True)
    else:
        print(f"Test project kept for inspection: {root}")
    sys.exit(0 if ok else 1)


if __name__ == "__main__":
    main()
